"""
批量增益测试模块
读取Excel配置文件，遍历测试不同增益控制字的实际增益

Excel配置文件格式：
    LNA增益 | TIA增益 | BBF增益 | PGA增益 | 信号源功率(dBm)
    8       | 1       | 3       | 6       | -80
"""
import pyvisa
from logic import serial_lib
import time
import numpy as np
import re
import serial
from openpyxl import load_workbook, Workbook
from pages.layout import log
from logic.iq_analyzer import analyze_iq_signal
from datetime import datetime
import os


class GainSweepError(Exception):
    """增益扫描相关错误"""
    pass


def build_gain_word(lna: int, tia: int, bbf: int, pga: int) -> int:
    """
    根据各级增益构建增益控制字

    参数:
        lna: LNA增益 (0-15)
        tia: TIA增益 (0-15)
        bbf: BBF增益 (0-15)
        pga: PGA增益 (0-15)

    返回:
        16位增益控制字
    """
    # 假设格式: [LNA:4bit][TIA:4bit][BBF:4bit][PGA:4bit]
    gain_word = ((lna & 0xF) << 12) | ((tia & 0xF) << 8) | ((bbf & 0xF) << 4) | (pga & 0xF)
    return gain_word


def parse_gain_word(gain_word: int) -> tuple:
    """
    解析增益控制字为各级增益

    返回: (lna, tia, bbf, pga)
    """
    lna = (gain_word >> 12) & 0xF
    tia = (gain_word >> 8) & 0xF
    bbf = (gain_word >> 4) & 0xF
    pga = gain_word & 0xF
    return lna, tia, bbf, pga


def load_gain_config(file_path: str) -> list:
    """
    从Excel文件加载增益配置

    Excel格式（从第2行开始，第1行为表头）：
    A列: LNA增益 (0-15)
    B列: TIA增益 (0-15)
    C列: BBF增益 (0-15)
    D列: PGA增益 (0-15)
    E列: 信号源功率 (dBm)

    返回: [(lna, tia, bbf, pga, signal_power, gain_word), ...]
    """
    if not os.path.exists(file_path):
        raise GainSweepError(f"配置文件不存在: {file_path}")

    wb = None
    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active

        configs = []
        for row in ws.iter_rows(min_row=2, values_only=True):  # 跳过表头
            if row[0] is None:
                continue

            # 解析各级增益
            lna = int(row[0]) if row[0] is not None else 0
            tia = int(row[1]) if row[1] is not None else 0
            bbf = int(row[2]) if row[2] is not None else 0
            pga = int(row[3]) if row[3] is not None else 0

            # 解析信号源功率
            signal_power = float(row[4]) if len(row) > 4 and row[4] is not None else -80.0

            # 构建增益控制字
            gain_word = build_gain_word(lna, tia, bbf, pga)

            configs.append((lna, tia, bbf, pga, signal_power, gain_word))

        log(f"成功加载 {len(configs)} 个增益配置", color="green")
        return configs

    except Exception as e:
        raise GainSweepError(f"读取配置文件失败: {e}")
    finally:
        if wb is not None:
            wb.close()


def hex2sint_list(hex_string_list, bits):
    """十六进制字符串列表转有符号整数列表"""
    result_list = []
    for hex_string in hex_string_list:
        hex_value = int(hex_string, 16) & (2**bits - 1)
        if hex_value > (2**(bits-1) - 1):
            decimal_value = hex_value - 2**bits
        else:
            decimal_value = hex_value
        result_list.append(decimal_value)
    return result_list


def gain_sweep(
    config_file: str,
    ip_str: str,
    com_num: int,
    buand: int = 115200,
    fs: float = 24e6,
    SN: int = 4096,
    vpp: float = 1.1,
    nbit: int = 12,
    cable_loss: float = 0.65,
    tone_freq: float = 1.0,
    chn: int = 19,
    ble_mode: str = "LE1M",
    output_file: str = None,
    progress_callback=None,
    chart_update_callback=None,
    result_callback=None,
    stop_flag=None
) -> list:
    """
    批量增益扫描测试

    参数:
        config_file: Excel配置文件路径
        ip_str: 信号源IP地址
        com_num: 串口号
        buand: 波特率
        fs: 采样率
        SN: 采样点数
        vpp: ADC参考电压
        nbit: ADC位数
        cable_loss: 线损
        tone_freq: 单音频偏(MHz) - 已废弃，实际使用中频（LE1M/LES2/LES8=1MHz, LE2M=2MHz）
        chn: 测试信道
        ble_mode: BLE制式
        output_file: 输出结果文件路径
        progress_callback: 进度回调函数 (current, total, message)
        chart_update_callback: 图表更新回调函数
        result_callback: 结果回调函数 (lna, tia, bbf, pga, measured_gain)
        stop_flag: 停止标志，callable返回True时停止测试

    返回:
        测试结果列表 [(lna, tia, bbf, pga, signal_power, gain_word, measured_gain, signal_dBm, noise_dBm), ...]

    说明:
        接收机为低中频架构：
        - LE1M/LES2/LES8: Lo比信道低1MHz，中频=1MHz
        - LE2M: Lo比信道低2MHz，中频=2MHz
        信号源频率设置为信道实际频率，单音功率在中频处测量
    """
    # 加载配置
    configs = load_gain_config(config_file)
    if not configs:
        raise GainSweepError("配置文件为空")

    total = len(configs)
    results = []

    # BLE速率映射
    ble_rate_map = {'LE1M': 0, 'LE2M': 1, 'LES2': 2, 'LES8': 3}
    rate_num = ble_rate_map.get(ble_mode, 0)

    # BLE中频偏移映射（低中频架构）
    ble_if_offset_map = {'LE1M': 1, 'LE2M': 2, 'LES2': 1, 'LES8': 1}
    if_offset_mhz = ble_if_offset_map.get(ble_mode, 1)

    # 连接信号源
    try:
        rm = pyvisa.ResourceManager()
        N5182B = rm.open_resource(f'TCPIP0::{ip_str}::INSTR')
        log(f"成功连接信号源: {ip_str}", color="green")
    except pyvisa.VisaIOError as e:
        log(f"连接信号源失败: {e}", color="red")
        raise GainSweepError(f"连接信号源失败: {e}")

    # 连接串口
    try:
        SerialPort = serial_lib.SerialCommunication(com_num, buand)
        log("打开串口成功", color="green")
        SerialPort.write_cmd('echoclose 0')
    except serial.SerialException as se:
        N5182B.close()
        rm.close()
        se_str = str(se)
        if "PermissionError" in se_str or "Access is denied" in se_str:
            msg = f"COM{com_num} 被占用"
        elif "FileNotFoundError" in se_str:
            msg = f"未发现 COM{com_num}"
        else:
            msg = f"串口错误: {se}"
        raise GainSweepError(msg)

    try:
        # 配置信号源为单音模式
        N5182B.write(':SOURce:RADio:ARB:STATe OFF')
        N5182B.write(':OUTPut:MODulation:STATe OFF')

        # 设置频率为信道实际频率（不加偏移）
        freq_mhz = 2402 + chn * 2
        N5182B.write(f':FREQuency:FIXed {freq_mhz} MHz')
        log(f"信号源频率: {freq_mhz} MHz, BLE制式: {ble_mode}, 中频: {if_offset_mhz} MHz", color="blue")

        for idx, (lna, tia, bbf, pga, signal_power, gain_word) in enumerate(configs):
            # 检查停止标志
            if stop_flag and callable(stop_flag) and stop_flag():
                log("用户请求停止测试", color="yellow")
                break

            if progress_callback:
                progress_callback(idx + 1, total, f"测试: LNA={lna} TIA={tia} BBF={bbf} PGA={pga}")

            log(f"[{idx+1}/{total}] LNA={lna}, TIA={tia}, BBF={bbf}, PGA={pga} (0x{gain_word:04X}), 功率={signal_power}dBm")

            # 设置信号源功率
            N5182B.write(f':POWer:LEVel {signal_power + cable_loss} dBm')
            N5182B.write(':OUTPut:STATe ON')
            time.sleep(0.1)

            # 启动BLE接收
            SerialPort.write_cmd(f'amtBleRxStartInf {rate_num} {chn + 1}')

            # 设置增益控制字到寄存器
            # 0x20203300: 增益控制字寄存器
            # 0x20203390: 强制使用寄存器控制字（写1启用）
            SerialPort.write_reg(0x20203300, gain_word)
            SerialPort.write_reg(0x20203390, 0x1)

            # 执行IQ dump
            SerialPort.write_cmd('amtbleiqdump 0 0 0x7fff 0x3a98')
            time.sleep(0.1)

            # 触发采集
            SerialPort.write_reg(0x2020e000, 0x0)
            N5182B.write(':OUTPut:STATe OFF')

            # 关闭强制增益控制
            SerialPort.write_reg(0x20203390, 0x0)

            # 停止接收
            SerialPort.write_cmd('amtBleRxStop')
            SerialPort.flush_Input()

            # 读取IQ数据
            SerialPort.write_cmd(f'read 0x20000000 {SN * 4}')
            temp = SerialPort.read_all_bytes()
            rxdump_data = temp.decode('UTF-8')

            # 解析数据
            hex_list = re.findall(r': ([0-9a-fA-F]+)\r\n', rxdump_data)
            if not hex_list:
                log(f"  警告: 未读取到有效数据", color="yellow")
                results.append((lna, tia, bbf, pga, signal_power, gain_word, None, None, None))
                continue

            Idata = [hex_value[2:5] for hex_value in hex_list]
            Qdata = [hex_value[5:8] for hex_value in hex_list]
            Idata_dec = hex2sint_list(Idata, 12)
            Qdata_dec = hex2sint_list(Qdata, 12)

            # 分析IQ信号
            # 目标单音频率 = 中频（1MHz或2MHz），而不是用户设置的tone_freq
            analysis = analyze_iq_signal(
                Idata=Idata_dec,
                Qdata=Qdata_dec,
                fs=fs, vpp=vpp, nbit=nbit,
                sg_pwr=signal_power,
                dump_noise=0,  # tone模式
                figure_off=0,
                target_tone_freq=if_offset_mhz * 1e6,  # 使用中频作为目标单音频率
                chn=chn,
                ble_mode=ble_mode
            )

            # 更新图表
            if chart_update_callback and 'time_chart' in analysis and 'freq_chart' in analysis:
                try:
                    chart_update_callback(analysis['time_chart'], analysis['freq_chart'])
                except Exception as e:
                    log(f"更新图表失败: {e}", color="yellow")

            measured_gain = analysis.get('gain')
            signal_dBm = analysis.get('signal')
            noise_dBm = analysis.get('noise')

            if measured_gain is not None:
                log(f"  测量增益: {measured_gain:.2f} dB, 信号: {signal_dBm:.2f} dBm", color="green")
            else:
                log(f"  测量失败", color="red")

            results.append((lna, tia, bbf, pga, signal_power, gain_word, measured_gain, signal_dBm, noise_dBm))

            # 调用结果回调，实时更新界面表格
            if result_callback:
                try:
                    result_callback(lna, tia, bbf, pga, measured_gain)
                except Exception as e:
                    log(f"更新结果表格失败: {e}", color="yellow")

            time.sleep(0.05)  # 短暂延迟

    finally:
        # 清理资源
        SerialPort.close()
        N5182B.close()
        rm.close()
        log("测试完成，资源已释放", color="green")

    # 保存结果到Excel
    if output_file:
        save_results_to_excel(results, output_file, ble_mode, chn, if_offset_mhz)

    return results


def save_results_to_excel(results: list, output_file: str, ble_mode: str, chn: int, if_offset_mhz: float):
    """保存测试结果到Excel文件"""
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "增益测试结果"

        # 写入测试信息
        ws['A1'] = "批量增益测试结果"
        ws['A2'] = f"测试时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws['A3'] = f"BLE制式: {ble_mode}, 信道: {chn}, 中频: {if_offset_mhz} MHz"

        # 表头
        headers = ['LNA', 'TIA', 'BBF', 'PGA', '增益控制字(Hex)', '信号源功率(dBm)',
                   '测量增益(dB)', '信号功率(dBm)', '噪声功率(dBm)']
        for col, header in enumerate(headers, 1):
            ws.cell(row=5, column=col, value=header)

        # 数据
        for row_idx, (lna, tia, bbf, pga, signal_power, gain_word, measured_gain, signal_dBm, noise_dBm) in enumerate(results, 6):
            ws.cell(row=row_idx, column=1, value=lna)
            ws.cell(row=row_idx, column=2, value=tia)
            ws.cell(row=row_idx, column=3, value=bbf)
            ws.cell(row=row_idx, column=4, value=pga)
            ws.cell(row=row_idx, column=5, value=f"0x{gain_word:04X}")
            ws.cell(row=row_idx, column=6, value=signal_power)
            ws.cell(row=row_idx, column=7, value=measured_gain)
            ws.cell(row=row_idx, column=8, value=signal_dBm)
            ws.cell(row=row_idx, column=9, value=noise_dBm)

        # 调整列宽
        column_widths = [6, 6, 6, 6, 16, 16, 14, 14, 14]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + col)].width = width

        wb.save(output_file)
        log(f"结果已保存到: {output_file}", color="green")

    except Exception as e:
        log(f"保存结果失败: {e}", color="red")


# 各级增益参数 (基准增益 + 每档步进)
# 增益 = 基准增益 + 控制字 * 步进
# 这些值需要根据实际芯片规格调整
DEFAULT_LNA_BASE_DB = 0   # LNA控制字0时的增益(dB)
DEFAULT_LNA_STEP_DB = 3   # LNA每档增益步进(dB)
DEFAULT_TIA_BASE_DB = 0   # TIA控制字0时的增益(dB)
DEFAULT_TIA_STEP_DB = 3   # TIA每档增益步进(dB)
DEFAULT_BBF_BASE_DB = 0   # BBF控制字0时的增益(dB)
DEFAULT_BBF_STEP_DB = 2   # BBF每档增益步进(dB)
DEFAULT_PGA_BASE_DB = 0   # PGA控制字0时的增益(dB)
DEFAULT_PGA_STEP_DB = 2   # PGA每档增益步进(dB)


def estimate_gain_for_stage(ctrl_word: int, base_db: float, step_db: float) -> float:
    """
    根据控制字计算该级增益
    增益 = 基准增益 + 控制字 * 步进
    """
    return base_db + ctrl_word * step_db


def estimate_total_gain(lna: int, tia: int, bbf: int, pga: int,
                        lna_base_db: float = DEFAULT_LNA_BASE_DB,
                        lna_step_db: float = DEFAULT_LNA_STEP_DB,
                        tia_base_db: float = DEFAULT_TIA_BASE_DB,
                        tia_step_db: float = DEFAULT_TIA_STEP_DB,
                        bbf_base_db: float = DEFAULT_BBF_BASE_DB,
                        bbf_step_db: float = DEFAULT_BBF_STEP_DB,
                        pga_base_db: float = DEFAULT_PGA_BASE_DB,
                        pga_step_db: float = DEFAULT_PGA_STEP_DB) -> float:
    """
    根据各级控制字估算总增益

    参数:
        lna, tia, bbf, pga: 各级增益控制字 (0-15)
        xxx_base_db: 该级控制字0时的基准增益(dB)
        xxx_step_db: 该级每档增益步进(dB)

    返回:
        估算的总增益 (dB)
    """
    lna_gain = estimate_gain_for_stage(lna, lna_base_db, lna_step_db)
    tia_gain = estimate_gain_for_stage(tia, tia_base_db, tia_step_db)
    bbf_gain = estimate_gain_for_stage(bbf, bbf_base_db, bbf_step_db)
    pga_gain = estimate_gain_for_stage(pga, pga_base_db, pga_step_db)
    return lna_gain + tia_gain + bbf_gain + pga_gain


def generate_gain_configs(
    lna_range: tuple = (0, 8),
    tia_range: tuple = (0, 8),
    bbf_range: tuple = (0, 8),
    pga_range: tuple = (0, 8),
    lna_base_db: float = DEFAULT_LNA_BASE_DB,
    lna_step_db: float = DEFAULT_LNA_STEP_DB,
    tia_base_db: float = DEFAULT_TIA_BASE_DB,
    tia_step_db: float = DEFAULT_TIA_STEP_DB,
    bbf_base_db: float = DEFAULT_BBF_BASE_DB,
    bbf_step_db: float = DEFAULT_BBF_STEP_DB,
    pga_base_db: float = DEFAULT_PGA_BASE_DB,
    pga_step_db: float = DEFAULT_PGA_STEP_DB,
    target_output_dbm: float = -5.0
) -> list:
    """
    根据范围参数生成所有增益配置组合（控制字步进为1，逐个遍历）

    参数:
        xxx_range: 控制字范围 (min, max)
        xxx_base_db: 控制字0时的基准增益(dB)
        xxx_step_db: 每档增益步进(dB)
        target_output_dbm: 目标输出功率 (dBm)，信号源功率 = 目标输出 - 估算增益

    返回:
        [(lna, tia, bbf, pga, signal_power), ...]
    """
    configs = []

    # 控制字步进固定为1，逐个遍历所有档位
    lna_values = list(range(lna_range[0], lna_range[1] + 1))
    tia_values = list(range(tia_range[0], tia_range[1] + 1))
    bbf_values = list(range(bbf_range[0], bbf_range[1] + 1))
    pga_values = list(range(pga_range[0], pga_range[1] + 1))

    # 遍历所有组合
    for lna in lna_values:
        for tia in tia_values:
            for bbf in bbf_values:
                for pga in pga_values:
                    # 估算总增益
                    total_gain = estimate_total_gain(
                        lna, tia, bbf, pga,
                        lna_base_db, lna_step_db,
                        tia_base_db, tia_step_db,
                        bbf_base_db, bbf_step_db,
                        pga_base_db, pga_step_db
                    )
                    # 计算信号源功率: 信号源功率 + 增益 = 目标输出
                    signal_power = target_output_dbm - total_gain
                    configs.append((lna, tia, bbf, pga, signal_power))

    return configs


def create_sample_config(file_path: str):
    """创建示例配置文件（使用默认参数）"""
    create_auto_config(
        file_path=file_path,
        lna_range=(0, 8),
        tia_range=(0, 4),
        bbf_range=(0, 4),
        pga_range=(0, 4),
        target_output_dbm=-5.0
    )


def create_auto_config(
    file_path: str,
    lna_range: tuple = (0, 8),
    tia_range: tuple = (0, 8),
    bbf_range: tuple = (0, 8),
    pga_range: tuple = (0, 8),
    lna_base_db: float = DEFAULT_LNA_BASE_DB,
    lna_step_db: float = DEFAULT_LNA_STEP_DB,
    tia_base_db: float = DEFAULT_TIA_BASE_DB,
    tia_step_db: float = DEFAULT_TIA_STEP_DB,
    bbf_base_db: float = DEFAULT_BBF_BASE_DB,
    bbf_step_db: float = DEFAULT_BBF_STEP_DB,
    pga_base_db: float = DEFAULT_PGA_BASE_DB,
    pga_step_db: float = DEFAULT_PGA_STEP_DB,
    target_output_dbm: float = -5.0
):
    """
    自动生成增益测试配置文件

    参数:
        file_path: 输出文件路径
        xxx_range: 控制字范围 (min, max)
        xxx_base_db: 控制字0时的基准增益(dB)
        xxx_step_db: 每档增益步进(dB)
        target_output_dbm: 目标输出功率 (dBm)
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "增益配置"

    # 表头
    ws['A1'] = "LNA"
    ws['B1'] = "TIA"
    ws['C1'] = "BBF"
    ws['D1'] = "PGA"
    ws['E1'] = "信号源功率(dBm)"
    ws['F1'] = "估算增益(dB)"

    # 生成配置
    configs = generate_gain_configs(
        lna_range=lna_range,
        tia_range=tia_range,
        bbf_range=bbf_range,
        pga_range=pga_range,
        lna_base_db=lna_base_db,
        lna_step_db=lna_step_db,
        tia_base_db=tia_base_db,
        tia_step_db=tia_step_db,
        bbf_base_db=bbf_base_db,
        bbf_step_db=bbf_step_db,
        pga_base_db=pga_base_db,
        pga_step_db=pga_step_db,
        target_output_dbm=target_output_dbm
    )

    # 写入数据
    for row_idx, (lna, tia, bbf, pga, signal_power) in enumerate(configs, 2):
        ws.cell(row=row_idx, column=1, value=lna)
        ws.cell(row=row_idx, column=2, value=tia)
        ws.cell(row=row_idx, column=3, value=bbf)
        ws.cell(row=row_idx, column=4, value=pga)
        ws.cell(row=row_idx, column=5, value=round(signal_power, 1))
        # 估算增益列（方便查看）
        estimated_gain = estimate_total_gain(
            lna, tia, bbf, pga,
            lna_base_db, lna_step_db,
            tia_base_db, tia_step_db,
            bbf_base_db, bbf_step_db,
            pga_base_db, pga_step_db
        )
        ws.cell(row=row_idx, column=6, value=estimated_gain)

    # 调整列宽
    for col in ['A', 'B', 'C', 'D']:
        ws.column_dimensions[col].width = 8
    ws.column_dimensions['E'].width = 18
    ws.column_dimensions['F'].width = 14

    wb.save(file_path)
    log(f"自动配置文件已创建: {file_path}, 共 {len(configs)} 个配置", color="green")
